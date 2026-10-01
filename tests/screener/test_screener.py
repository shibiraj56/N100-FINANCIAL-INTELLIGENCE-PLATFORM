"""Unit tests for the Screener Engine (Epic 03, Day 21)."""

import os
import openpyxl
import pandas as pd
import pytest
from src.screener.engine import ScreenerEngine
from src.screener.exporter import export_screener_workbook


@pytest.fixture(scope="module")
def screener():
    return ScreenerEngine()


def test_screener_config_loading(screener):
    """Verify screener config loads correctly with all 6 presets."""
    assert screener.config is not None
    presets = screener.config.get("presets", {})
    expected_presets = [
        "quality_compounder",
        "value_pick",
        "growth_accelerator",
        "dividend_champion",
        "debt_free_blue_chip",
        "turnaround_watch",
    ]
    for p in expected_presets:
        assert p in presets, f"Missing preset: {p}"


def test_universe_data_fetch(screener):
    """Verify screener fetches unified dataset for all 92 companies."""
    df = screener.get_universe_data(year="2024-03")
    assert not df.empty
    assert len(df) == 92
    expected_cols = [
        "company_id", "company_name", "broad_sector", "sales",
        "return_on_equity_pct", "debt_to_equity", "free_cash_flow_cr",
        "pe_ratio", "pb_ratio", "dividend_yield_pct", "composite_quality_score",
        "sector_composite_score"
    ]
    for col in expected_cols:
        assert col in df.columns, f"Missing column {col} in universe data"


def test_preset_quality_compounder(screener):
    """Quality Compounder preset test: returns 5-50 companies meeting criteria."""
    df = screener.run_quality_compounder(year="2024-03")
    assert 5 <= len(df) <= 50, f"Expected 5-50 companies, got {len(df)}"
    assert (df["return_on_equity_pct"] >= 15.0).all()
    assert (df["free_cash_flow_cr"] > 0).all()
    assert (df["revenue_cagr_5yr"] >= 10.0).all()

    # Verify Financials sector carve-out for D/E
    non_fin = df[df["broad_sector"].astype(str).str.lower() != "financials"]
    assert (non_fin["debt_to_equity"] <= 1.0).all()


def test_preset_value_pick(screener):
    """Value Pick preset test: returns 5-50 companies meeting value criteria."""
    df = screener.run_value_pick(year="2024-03")
    assert 5 <= len(df) <= 50, f"Expected 5-50 companies, got {len(df)}"
    assert (df["pe_ratio"] <= 25.0).all()
    assert (df["pb_ratio"] <= 5.0).all()
    assert (df["dividend_yield_pct"] >= 1.0).all()


def test_preset_growth_accelerator(screener):
    """Growth Accelerator preset test: returns 5-50 companies."""
    df = screener.run_growth_accelerator(year="2024-03")
    assert 5 <= len(df) <= 50, f"Expected 5-50 companies, got {len(df)}"
    assert (df["pat_cagr_5yr"] >= 20.0).all()
    assert (df["revenue_cagr_5yr"] >= 15.0).all()


def test_preset_dividend_champion(screener):
    """Dividend Champion preset test: returns 5-50 companies."""
    df = screener.run_dividend_champion(year="2024-03")
    assert 5 <= len(df) <= 50, f"Expected 5-50 companies, got {len(df)}"
    assert (df["dividend_yield_pct"] >= 2.0).all()
    assert (df["dividend_payout_ratio_pct"] <= 80.0).all()
    assert (df["free_cash_flow_cr"] > 0).all()


def test_preset_debt_free_blue_chip(screener):
    """Debt Free Blue Chip preset test: returns 5-50 companies strictly debt-free."""
    df = screener.run_debt_free_blue_chip(year="2024-03")
    assert 5 <= len(df) <= 50, f"Expected 5-50 companies, got {len(df)}"
    assert (df["debt_to_equity"] == 0.0).all()
    assert (df["return_on_equity_pct"] >= 12.0).all()
    assert (df["sales"] >= 5000.0).all()


def test_preset_turnaround_watch(screener):
    """Turnaround Watch preset test: returns 5-50 companies."""
    df = screener.run_turnaround_watch(year="2024-03")
    assert 5 <= len(df) <= 50, f"Expected 5-50 companies, got {len(df)}"
    assert (df["revenue_cagr_3yr"] >= 10.0).all()
    assert (df["free_cash_flow_cr"] > 0).all()


def test_financials_de_carve_out(screener):
    """Verify that Financials companies are not excluded by de_max."""
    df_uni = screener.get_universe_data(year="2024-03")
    filtered = screener.apply_filters(df_uni, {"de_max": 0.01})
    financials_in_res = filtered[filtered["broad_sector"].astype(str).str.lower() == "financials"]
    assert not financials_in_res.empty, "Financials should be preserved under D/E max filter"


def test_icr_debt_free_infinity(screener):
    """Verify that companies labeled Debt Free or with D/E=0 pass high ICR filters."""
    # Test with universe data
    df_uni = screener.get_universe_data(year="2024-03")
    filtered = screener.apply_filters(df_uni, {"icr_min": 100.0})
    debt_free_comps = filtered[filtered["debt_to_equity"] == 0.0]
    assert not debt_free_comps.empty, "Debt Free companies should pass any ICR min"

    # Test with explicit 'Debt Free' label in mock data
    mock_df = pd.DataFrame([
        {"company_id": "CO1", "interest_coverage": None, "icr_label": "Debt Free", "debt_to_equity": 0.0, "broad_sector": "IT"},
        {"company_id": "CO2", "interest_coverage": 1.2, "icr_label": "1.2x", "debt_to_equity": 1.5, "broad_sector": "Materials"},
    ])
    res = screener.apply_filters(mock_df, {"icr_min": 10.0})
    assert len(res) == 1
    assert res.iloc[0]["company_id"] == "CO1"


def test_sector_relative_scores(screener):
    """Verify sector-relative composite score computation."""
    df_uni = screener.get_universe_data(year="2024-03")
    assert "sector_composite_score" in df_uni.columns
    scores = df_uni["sector_composite_score"].dropna()
    assert (scores >= 0.0).all()
    assert (scores <= 100.0).all()


def test_custom_filter_combination(screener):
    """Test custom multi-metric filtering."""
    df_uni = screener.get_universe_data(year="2024-03")
    custom_filters = {
        "roe_min": 20.0,
        "opm_min": 15.0,
        "net_profit_min": 1000.0,
    }
    filtered = screener.apply_filters(df_uni, custom_filters)
    assert not filtered.empty
    assert (filtered["return_on_equity_pct"] >= 20.0).all()
    assert (filtered["operating_profit_margin_pct"] >= 15.0).all()
    assert (filtered["net_profit"] >= 1000.0).all()


def test_screener_excel_export_structure():
    """Verify output/screener_output.xlsx exists and has 6 valid sheets."""
    wb_path = os.path.join("output", "screener_output.xlsx")
    assert os.path.exists(wb_path), f"{wb_path} does not exist"
    wb = openpyxl.load_workbook(wb_path)
    assert len(wb.sheetnames) == 6
    expected_sheets = [
        "Quality Compounders",
        "Value Picks",
        "Growth Accelerators",
        "Dividend Champions",
        "Debt-Free Blue Chips",
        "Turnaround Watch",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames
        ws = wb[s]
        assert ws.max_row >= 2  # At least header + 1 row
        assert ws.max_column == 20  # 20 KPI columns

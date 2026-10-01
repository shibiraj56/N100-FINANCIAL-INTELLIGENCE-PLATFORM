"""Unit tests for Valuation Engine and Overvaluation Flags (Epic 06, Day 26)."""

import os
import openpyxl
import pandas as pd
import pytest

from src.analytics.valuation import (
    ValuationEngine,
    calculate_fcf_yield,
    classify_valuation_flag,
)


def test_calculate_fcf_yield():
    """Verify FCF yield calculation."""
    assert calculate_fcf_yield(1000.0, 20000.0) == 5.0
    assert calculate_fcf_yield(-500.0, 10000.0) == -5.0
    assert calculate_fcf_yield(None, 10000.0) is None
    assert calculate_fcf_yield(500.0, 0.0) is None
    assert calculate_fcf_yield(500.0, -100.0) is None


def test_classify_valuation_flag():
    """Verify valuation classification thresholds."""
    # Sector median = 20.0
    # Caution threshold = > 30.0 (1.5x)
    # Discount threshold = < 14.0 (0.7x)
    assert classify_valuation_flag(35.0, 20.0) == "Caution"
    assert classify_valuation_flag(12.0, 20.0) == "Discount"
    assert classify_valuation_flag(22.0, 20.0) == "Fair"
    assert classify_valuation_flag(20.0, 20.0) == "Fair"
    # Edge cases
    assert classify_valuation_flag(None, 20.0) == "Fair"
    assert classify_valuation_flag(-5.0, 20.0) == "Fair"


def test_valuation_dataset_structure():
    """Verify ValuationEngine computes full 92-company dataset with all required metrics."""
    engine = ValuationEngine()
    df_val = engine.compute_valuation_dataset(year="2024-03")
    assert len(df_val) == 92
    expected_cols = [
        "company_id", "company_name", "sector", "P/E", "P/B",
        "EV/EBITDA", "FCF_yield_pct", "5yr_median_PE", "PE_vs_sector_median_pct", "flag"
    ]
    for col in expected_cols:
        assert col in df_val.columns, f"Missing column: {col}"
    
    # Check flags are strictly among allowed values
    allowed_flags = {"Caution", "Discount", "Fair"}
    assert set(df_val["flag"].unique()).issubset(allowed_flags)


def test_valuation_summary_excel_artifact():
    """Verify output/valuation_summary.xlsx exists and has 92 company records."""
    summary_path = os.path.join("output", "valuation_summary.xlsx")
    assert os.path.exists(summary_path), f"{summary_path} does not exist"
    wb = openpyxl.load_workbook(summary_path)
    ws = wb.active
    # Header + 92 company rows = 93 rows
    assert ws.max_row == 93, f"Expected 93 rows (header + 92 data rows), got {ws.max_row}"
    assert ws.max_column == 10, f"Expected 10 columns, got {ws.max_column}"


def test_valuation_flags_csv_artifact():
    """Verify output/valuation_flags.csv exists and only contains Caution/Discount companies."""
    flags_path = os.path.join("output", "valuation_flags.csv")
    assert os.path.exists(flags_path), f"{flags_path} does not exist"
    df_flags = pd.read_csv(flags_path)
    assert not df_flags.empty
    assert set(df_flags["flag"].unique()).issubset({"Caution", "Discount"})
    assert "Fair" not in df_flags["flag"].values

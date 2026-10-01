"""Unit tests for Peer Comparison & Intra-Group Percentile Engine (Epic 04, Day 21)."""

import os
import sqlite3
import openpyxl
import pandas as pd
import pytest
from src.analytics.peer import PeerEngine, compute_percent_rank


@pytest.fixture(scope="module")
def peer_engine():
    engine = PeerEngine()
    conn = sqlite3.connect(engine.db_path)
    cur = conn.cursor()
    cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='peer_percentiles'")
    tbl = cur.fetchone()
    count = 0
    if tbl:
        cur.execute("SELECT COUNT(*) FROM peer_percentiles")
        count = cur.fetchone()[0]
    conn.close()
    if count == 0:
        engine.compute_peer_percentiles()
    return engine


def test_compute_percent_rank_standard():
    """Verify standard ascending PERCENT_RANK calculation."""
    vals = [10.0, 20.0, 30.0]
    ranks = compute_percent_rank(vals, is_inverse=False)
    assert ranks == [0.0, 0.5, 1.0]


def test_compute_percent_rank_inverse():
    """Verify inverse descending PERCENT_RANK calculation (e.g. for D/E)."""
    vals = [10.0, 20.0, 30.0]
    ranks = compute_percent_rank(vals, is_inverse=True)
    # Lowest value (10.0) should get rank 1.0, highest (30.0) should get 0.0
    assert ranks == [1.0, 0.5, 0.0]


def test_peer_percentiles_row_count(peer_engine):
    """Verify peer_percentiles table has exactly 560 rows (56 companies * 10 metrics)."""
    conn = sqlite3.connect(peer_engine.db_path)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM peer_percentiles WHERE year = '2024-03'")
    count = cur.fetchone()[0]
    conn.close()
    assert count == 560, f"Expected 560 rows in peer_percentiles, got {count}"


def test_peer_groups_coverage(peer_engine):
    """Verify all 11 peer groups are represented in peer_percentiles."""
    conn = sqlite3.connect(peer_engine.db_path)
    df = pd.read_sql("SELECT DISTINCT peer_group_name FROM peer_percentiles", conn)
    conn.close()
    assert len(df) == 11, f"Expected 11 peer groups, found {len(df)}"


def test_it_services_tcs_roe_benchmark(peer_engine):
    """Spot-check: TCS in IT Services must have highest ROE and percentile_rank = 1.0."""
    conn = sqlite3.connect(peer_engine.db_path)
    query = """
    SELECT company_id, metric, value, percentile_rank 
    FROM peer_percentiles 
    WHERE peer_group_name = 'IT Services' AND metric = 'roe'
    ORDER BY percentile_rank DESC
    """
    df = pd.read_sql(query, conn)
    conn.close()
    assert not df.empty
    top_co = df.iloc[0]
    assert top_co["company_id"] == "TCS"
    assert top_co["percentile_rank"] == 1.0


def test_de_inverse_ranking(peer_engine):
    """Verify Debt-to-Equity is inversely ranked (lowest D/E receives highest percentile rank)."""
    conn = sqlite3.connect(peer_engine.db_path)
    query = """
    SELECT company_id, value, percentile_rank 
    FROM peer_percentiles 
    WHERE peer_group_name = 'Automobiles' AND metric = 'de'
    ORDER BY value ASC
    """
    df = pd.read_sql(query, conn)
    conn.close()
    assert not df.empty
    # The lowest D/E company should have the highest percentile rank
    lowest_de = df.iloc[0]
    highest_de = df.iloc[-1]
    assert lowest_de["percentile_rank"] > highest_de["percentile_rank"]


def test_unassigned_company_handling(peer_engine):
    """Verify companies without a peer group return a friendly message without throwing an error."""
    res = peer_engine.get_company_percentiles("ABB", year="2024-03")
    assert isinstance(res, dict)
    assert res.get("message") == "No peer group assigned"
    assert res.get("company_id") == "ABB"


def test_percentile_rank_bounds(peer_engine):
    """Verify all percentile ranks are strictly within [0.0, 1.0]."""
    conn = sqlite3.connect(peer_engine.db_path)
    df = pd.read_sql("SELECT percentile_rank FROM peer_percentiles", conn)
    conn.close()
    ranks = df["percentile_rank"]
    assert (ranks >= 0.0).all()
    assert (ranks <= 1.0).all()


def test_peer_comparison_excel_sheets():
    """Verify output/peer_comparison.xlsx exists and has exactly 11 sheets."""
    excel_path = os.path.join("output", "peer_comparison.xlsx")
    assert os.path.exists(excel_path), f"{excel_path} does not exist"
    wb = openpyxl.load_workbook(excel_path)
    assert len(wb.sheetnames) == 11, f"Expected 11 sheets, got {len(wb.sheetnames)}"


def test_peer_comparison_excel_content():
    """Verify each sheet has benchmark row and group median summary row."""
    excel_path = os.path.join("output", "peer_comparison.xlsx")
    wb = openpyxl.load_workbook(excel_path)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Should have header + data rows + median row >= 4 rows
        assert ws.max_row >= 4
        # Verify median row at the bottom
        last_row_title = ws.cell(row=ws.max_row, column=2).value
        assert last_row_title == "Group Median"


def test_radar_charts_generation():
    """Verify radar charts are generated in reports/radar_charts/."""
    chart_dir = os.path.join("reports", "radar_charts")
    assert os.path.exists(chart_dir), f"{chart_dir} does not exist"
    png_files = [f for f in os.listdir(chart_dir) if f.endswith(".png")]
    assert len(png_files) >= 56, f"Expected at least 56 radar charts, found {len(png_files)}"

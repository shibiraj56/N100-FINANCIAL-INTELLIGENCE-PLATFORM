"""Screener Excel Exporter with openpyxl styling and conditional threshold formatting.

Generates output/screener_output.xlsx with 6 sheets, one per preset, with 20 KPI columns,
sorted by composite score, styled with green/red threshold cell fills and auto column widths.
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
import sys
from typing import Any, Dict, List, Optional

# Ensure project root is in sys.path
root_dir = str(Path(__file__).resolve().parent.parent.parent)
if root_dir not in sys.path:
    sys.path.insert(0, root_dir)

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from src.screener.engine import ScreenerEngine

logger = logging.getLogger(__name__)

# Colors
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10)
PASS_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
PASS_FONT = Font(name="Calibri", size=10, color="274E13", bold=True)
FAIL_FILL = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
FAIL_FONT = Font(name="Calibri", size=10, color="783F04")

THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)

KPI_COLUMNS = [
    ("company_id", "Ticker"),
    ("company_name", "Company Name"),
    ("broad_sector", "Broad Sector"),
    ("composite_quality_score", "Quality Score"),
    ("sector_composite_score", "Sector Score"),
    ("return_on_equity_pct", "ROE (%)"),
    ("roce_pct", "ROCE (%)"),
    ("net_profit_margin_pct", "NPM (%)"),
    ("operating_profit_margin_pct", "OPM (%)"),
    ("debt_to_equity", "D/E"),
    ("interest_coverage", "ICR"),
    ("free_cash_flow_cr", "FCF (₹ Cr)"),
    ("revenue_cagr_5yr", "5Y Rev CAGR (%)"),
    ("pat_cagr_5yr", "5Y PAT CAGR (%)"),
    ("eps_cagr_5yr", "5Y EPS CAGR (%)"),
    ("pe_ratio", "P/E"),
    ("pb_ratio", "P/B"),
    ("dividend_yield_pct", "Div Yield (%)"),
    ("dividend_payout_ratio_pct", "Payout (%)"),
    ("sales", "Revenue (₹ Cr)"),
]


def export_screener_workbook(
    preset_results: Dict[str, pd.DataFrame],
    output_path: str = "output/screener_output.xlsx",
) -> str:
    """Create and save the formatted multi-sheet screener Excel workbook."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    preset_thresholds = {
        "quality_compounder": {
            "return_on_equity_pct": (lambda v: v is not None and v > 15.0),
            "debt_to_equity": (lambda v: v is not None and v < 1.0),
            "free_cash_flow_cr": (lambda v: v is not None and v > 0),
            "revenue_cagr_5yr": (lambda v: v is not None and v > 10.0),
        },
        "value_pick": {
            "pe_ratio": (lambda v: v is not None and 0 < v <= 25.0),
            "pb_ratio": (lambda v: v is not None and 0 < v <= 5.0),
            "debt_to_equity": (lambda v: v is not None and v < 2.0),
            "dividend_yield_pct": (lambda v: v is not None and v >= 1.0),
        },
        "growth_accelerator": {
            "pat_cagr_5yr": (lambda v: v is not None and v > 20.0),
            "revenue_cagr_5yr": (lambda v: v is not None and v > 15.0),
            "debt_to_equity": (lambda v: v is not None and v < 2.0),
        },
        "dividend_champion": {
            "dividend_yield_pct": (lambda v: v is not None and v >= 2.0),
            "dividend_payout_ratio_pct": (lambda v: v is not None and v < 80.0),
            "free_cash_flow_cr": (lambda v: v is not None and v > 0),
        },
        "debt_free_blue_chip": {
            "debt_to_equity": (lambda v: v is not None and v == 0.0),
            "return_on_equity_pct": (lambda v: v is not None and v > 12.0),
            "sales": (lambda v: v is not None and v > 5000.0),
        },
        "turnaround_watch": {
            "revenue_cagr_3yr": (lambda v: v is not None and v > 10.0),
            "free_cash_flow_cr": (lambda v: v is not None and v > 0),
        },
    }

    sheet_names = {
        "quality_compounder": "Quality Compounders",
        "value_pick": "Value Picks",
        "growth_accelerator": "Growth Accelerators",
        "dividend_champion": "Dividend Champions",
        "debt_free_blue_chip": "Debt-Free Blue Chips",
        "turnaround_watch": "Turnaround Watch",
    }

    for p_key, df in preset_results.items():
        title = sheet_names.get(p_key, p_key[:31])
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True

        # Header Row
        headers = [disp for _, disp in KPI_COLUMNS]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER

        ws.row_dimensions[1].height = 25

        # Data Rows
        rule_map = preset_thresholds.get(p_key, {})

        for row_idx, row in enumerate(df.iterrows(), start=2):
            _, data = row
            is_financial = str(data.get("broad_sector", "")).lower() == "financials"

            for col_idx, (raw_col, _) in enumerate(KPI_COLUMNS, start=1):
                val = data.get(raw_col)

                # Format numbers nicely
                if isinstance(val, (float, int)):
                    if abs(val) > 1000 and "cagr" not in raw_col and "pct" not in raw_col and "ratio" not in raw_col:
                        cell_val = round(val, 1)
                    else:
                        cell_val = round(val, 2)
                elif pd.isna(val):
                    cell_val = "—"
                else:
                    cell_val = str(val)

                cell = ws.cell(row=row_idx, column=col_idx, value=cell_val)
                cell.font = ROW_FONT
                cell.border = THIN_BORDER

                # Alignment
                if col_idx in (1, 3):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 2:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

                # Conditional formatting
                if raw_col in rule_map:
                    cond_fn = rule_map[raw_col]
                    # Financials D/E exception
                    if raw_col == "debt_to_equity" and is_financial:
                        cell.fill = PASS_FILL
                        cell.font = PASS_FONT
                    elif cond_fn(val):
                        cell.fill = PASS_FILL
                        cell.font = PASS_FONT
                    else:
                        cell.fill = FAIL_FILL
                        cell.font = FAIL_FONT

            ws.row_dimensions[row_idx].height = 20

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

        # Freeze header
        ws.freeze_panes = "A2"

    wb.save(output_path)
    logger.info("Saved styled screener workbook to %s", output_path)
    return output_path


def generate_all_screener_outputs(
    db_path: str = "data/nifty100.db",
    config_path: str = "config/screener_config.yaml",
    output_path: str = "output/screener_output.xlsx",
) -> str:
    """Run all presets and export to styled Excel workbook."""
    engine = ScreenerEngine(db_path=db_path, config_path=config_path)
    results = engine.run_all_presets()
    return export_screener_workbook(results, output_path=output_path)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    path = generate_all_screener_outputs()
    print("Screener export generated successfully at:", path)

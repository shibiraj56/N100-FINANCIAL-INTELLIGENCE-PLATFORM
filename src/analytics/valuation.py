"""Valuation Module for N100 Companies (Epic 06, Day 26).

Computes FCF Yield, Sector Median P/E, 5-Year Median P/E, and Overvaluation/Discount flags.
Generates output/valuation_summary.xlsx and output/valuation_flags.csv.
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
import sqlite3
import sys
from typing import Any, Dict, List, Optional, Tuple

# Ensure project root is in sys.path
root_dir = str(Path(__file__).resolve().parent.parent.parent)
if root_dir not in sys.path:
    sys.path.insert(0, root_dir)

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

logger = logging.getLogger(__name__)

DEFAULT_DB_PATH = os.path.join(root_dir, "data", "nifty100.db")

# Excel Styling
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
CAUTION_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
CAUTION_FONT = Font(name="Calibri", size=10, bold=True, color="990000")
DISCOUNT_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
DISCOUNT_FONT = Font(name="Calibri", size=10, bold=True, color="274E13")
FAIR_FONT = Font(name="Calibri", size=10, color="000000")
DATA_FONT = Font(name="Calibri", size=10)

THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)


def calculate_fcf_yield(fcf_cr: Optional[float], market_cap_cr: Optional[float]) -> Optional[float]:
    """Compute Free Cash Flow Yield (%): (FCF / Market Cap) * 100."""
    if fcf_cr is None or market_cap_cr is None or market_cap_cr <= 0:
        return None
    return round((fcf_cr / market_cap_cr) * 100.0, 2)


def classify_valuation_flag(
    pe: Optional[float], sector_median_pe: Optional[float]
) -> str:
    """Classify company valuation multiple relative to sector median.
    
    - Caution: P/E > sector_median * 1.5
    - Discount: P/E < sector_median * 0.7
    - Fair: otherwise
    """
    if pe is None or pe <= 0 or sector_median_pe is None or sector_median_pe <= 0:
        return "Fair"
    if pe > sector_median_pe * 1.5:
        return "Caution"
    if pe < sector_median_pe * 0.7:
        return "Discount"
    return "Fair"


class ValuationEngine:
    """Computes comprehensive valuation metrics and generates audit reports."""

    def __init__(self, db_path: str = DEFAULT_DB_PATH):
        self.db_path = db_path

    def compute_valuation_dataset(self, year: str = "2024-03") -> pd.DataFrame:
        """Compute unified valuation metrics for all 92 companies for specified year."""
        conn = sqlite3.connect(self.db_path)

        # 1. Fetch latest valuation multiples and financials
        query_latest = """
        SELECT 
            m.company_id,
            c.company_name,
            COALESCE(s.broad_sector, 'Unknown') AS sector,
            m.pe_ratio AS pe,
            m.pb_ratio AS pb,
            m.ev_ebitda,
            m.market_cap_crore,
            m.dividend_yield_pct,
            r.free_cash_flow_cr
        FROM market_cap m
        JOIN companies c ON m.company_id = c.id
        LEFT JOIN sectors s ON m.company_id = s.company_id
        LEFT JOIN financial_ratios r ON m.company_id = r.company_id AND m.year = r.year
        WHERE m.year = ?
        """
        df_latest = pd.read_sql(query_latest, conn, params=(year,))

        # 2. Fetch historical P/E for 5-year median calculation
        query_hist = """
        SELECT company_id, year, pe_ratio 
        FROM market_cap 
        WHERE pe_ratio > 0
        ORDER BY year DESC
        """
        df_hist = pd.read_sql(query_hist, conn)
        conn.close()

        # Compute 5-year median P/E per company (taking up to 5 most recent years)
        five_yr_medians: Dict[str, float] = {}
        for cid, grp in df_hist.groupby("company_id"):
            recent_pe = grp.head(5)["pe_ratio"]
            if not recent_pe.empty:
                five_yr_medians[cid] = round(float(recent_pe.median()), 2)

        # Compute sector median P/E for the current year
        sector_medians: Dict[str, float] = {}
        for sec, grp in df_latest.groupby("sector"):
            valid_pe = grp[grp["pe"] > 0]["pe"]
            if not valid_pe.empty:
                sector_medians[sec] = round(float(valid_pe.median()), 2)
            else:
                sector_medians[sec] = 25.0

        # Build final dataset
        records: List[Dict[str, Any]] = []
        for _, row in df_latest.iterrows():
            cid = row["company_id"]
            cname = row["company_name"]
            sector = row["sector"]
            pe = round(float(row["pe"]), 2) if pd.notna(row["pe"]) else None
            pb = round(float(row["pb"]), 2) if pd.notna(row["pb"]) else None
            ev_ebitda = round(float(row["ev_ebitda"]), 2) if pd.notna(row["ev_ebitda"]) else None
            mc = row["market_cap_crore"]
            fcf = row["free_cash_flow_cr"]

            fcf_yield = calculate_fcf_yield(fcf, mc)
            five_yr_med = five_yr_medians.get(cid)
            sec_med = sector_medians.get(sector, 25.0)

            if pe is not None and pe > 0 and sec_med and sec_med > 0:
                pe_vs_sec_pct = round(((pe - sec_med) / sec_med) * 100.0, 2)
            else:
                pe_vs_sec_pct = None

            flag = classify_valuation_flag(pe, sec_med)

            records.append({
                "company_id": cid,
                "company_name": cname,
                "sector": sector,
                "P/E": pe,
                "P/B": pb,
                "EV/EBITDA": ev_ebitda,
                "FCF_yield_pct": fcf_yield,
                "5yr_median_PE": five_yr_med,
                "PE_vs_sector_median_pct": pe_vs_sec_pct,
                "flag": flag,
            })

        df_val = pd.DataFrame(records)
        df_val = df_val.sort_values(by=["sector", "company_id"]).reset_index(drop=True)
        return df_val

    def generate_reports(
        self,
        year: str = "2024-03",
        summary_excel_path: str = "output/valuation_summary.xlsx",
        flags_csv_path: str = "output/valuation_flags.csv",
    ) -> Tuple[str, str]:
        """Generate output/valuation_summary.xlsx and output/valuation_flags.csv."""
        df_val = self.compute_valuation_dataset(year=year)

        # 1. Export Excel
        os.makedirs(os.path.dirname(summary_excel_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Valuation Summary"

        headers = [
            "Ticker", "Company Name", "Sector", "P/E", "P/B",
            "EV/EBITDA", "FCF Yield (%)", "5Y Median P/E", "P/E vs Sector (%)", "Valuation Flag"
        ]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
        ws.row_dimensions[1].height = 25

        col_keys = [
            "company_id", "company_name", "sector", "P/E", "P/B",
            "EV/EBITDA", "FCF_yield_pct", "5yr_median_PE", "PE_vs_sector_median_pct", "flag"
        ]

        for r_idx, row in df_val.iterrows():
            row_num = r_idx + 2
            flag = row["flag"]
            for c_idx, key in enumerate(col_keys, start=1):
                val = row[key]
                cell = ws.cell(row=row_num, column=c_idx, value=val if val is not None else "—")
                cell.font = DATA_FONT
                cell.border = THIN_BORDER

                if c_idx in (1, 2, 3):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 10:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if flag == "Caution":
                        cell.fill = CAUTION_FILL
                        cell.font = CAUTION_FONT
                    elif flag == "Discount":
                        cell.fill = DISCOUNT_FILL
                        cell.font = DISCOUNT_FONT
                    else:
                        cell.font = FAIR_FONT
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

            ws.row_dimensions[row_num].height = 20

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        ws.freeze_panes = "A2"
        wb.save(summary_excel_path)
        logger.info("Saved valuation summary workbook to %s", summary_excel_path)

        # 2. Export Flags CSV (only Caution or Discount)
        os.makedirs(os.path.dirname(flags_csv_path), exist_ok=True)
        df_flags = df_val[df_val["flag"].isin(["Caution", "Discount"])].copy()
        df_flags.to_csv(flags_csv_path, index=False)
        logger.info("Saved valuation flags CSV to %s (%d companies flagged)", flags_csv_path, len(df_flags))

        return summary_excel_path, flags_csv_path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    engine = ValuationEngine()
    summary_path, flags_path = engine.generate_reports()
    print("Valuation summary generated at:", summary_path)
    print("Valuation flags CSV generated at:", flags_path)
